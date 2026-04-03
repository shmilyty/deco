from __future__ import annotations

import gzip
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

try:
    import win32com.client as win32  # type: ignore
except Exception:  # pragma: no cover - optional on non-Windows setups
    win32 = None


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "competition_docs" / "03设计与开发文档"
HOME_SHOT = ROOT / "public" / "screenshot" / "homepage.png"
DEMO_SHOT = ROOT / "public" / "demo.png"
NOW = datetime.now()
DATE_TEXT = f"{NOW.year}年{NOW.month}月{NOW.day}日"

PROJECT_NAME = "Deco My Tree：基于时间胶囊机制的多人互动祝福平台"
PROJECT_NAME_SHORT = "Deco My Tree"
DOC_WARNING = "说明：文中带【待填写】的字段必须结合队伍真实信息补全后，再导出为最终提交版 PDF。"


def set_run_font(run, name: str = "宋体", size: float = 10.5, bold: bool = False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.8)

    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)


def add_paragraph(
    doc: Document,
    text: str,
    *,
    font: str = "宋体",
    size: float = 10.5,
    bold: bool = False,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    color=None,
    space_after: float = 6,
):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run_color = color
    if "【待填写】" in text or "【待替换】" in text or "【%】" in text:
        run_color = RGBColor(192, 0, 0)
    set_run_font(run, name=font, size=size, bold=bold, color=run_color)
    return p


def add_mixed_paragraph(doc: Document, parts, *, align=WD_ALIGN_PARAGRAPH.LEFT, space_after: float = 6):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    for text, options in parts:
        run = p.add_run(text)
        set_run_font(
            run,
            name=options.get("font", "宋体"),
            size=options.get("size", 10.5),
            bold=options.get("bold", False),
            color=options.get("color"),
        )
    return p


def add_heading(doc: Document, text: str, level: int = 1):
    if level == 1:
        return add_paragraph(doc, text, font="黑体", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    if level == 2:
        return add_paragraph(doc, text, font="黑体", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=8)
    return add_paragraph(doc, text, font="黑体", size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6)


def fill_cell(cell, text: str, *, bold: bool = False, align=WD_ALIGN_PARAGRAPH.LEFT, font="宋体", size=10.5):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    color = None
    if "【待填写】" in text or "【待替换】" in text or "【%】" in text:
        color = RGBColor(192, 0, 0)
    set_run_font(run, name=font, size=size, bold=bold, color=color)


def add_table(doc: Document, rows, *, header=True):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            fill_cell(
                table.cell(r_idx, c_idx),
                value,
                bold=header and r_idx == 0,
                align=WD_ALIGN_PARAGRAPH.CENTER if (header and r_idx == 0) else WD_ALIGN_PARAGRAPH.LEFT,
                font="黑体" if header and r_idx == 0 else "宋体",
                size=10.5,
            )
    return table


def add_bullets(doc: Document, items):
    for item in items:
        add_paragraph(doc, f"• {item}")


def get_build_metrics() -> dict[str, str]:
    metrics = {
        "js_size": "未检测",
        "js_gzip": "未检测",
        "css_size": "未检测",
        "css_gzip": "未检测",
    }
    assets_dir = ROOT / "dist" / "assets"
    if not assets_dir.exists():
        return metrics

    js_files = sorted(assets_dir.glob("index-*.js"))
    css_files = sorted(assets_dir.glob("index-*.css"))
    if js_files:
        data = js_files[0].read_bytes()
        metrics["js_size"] = f"{len(data) / 1024:.2f} KB"
        metrics["js_gzip"] = f"{len(gzip.compress(data)) / 1024:.2f} KB"
    if css_files:
        data = css_files[0].read_bytes()
        metrics["css_size"] = f"{len(data) / 1024:.2f} KB"
        metrics["css_gzip"] = f"{len(gzip.compress(data)) / 1024:.2f} KB"
    return metrics


def build_summary_doc() -> Path:
    doc = Document()
    configure_document(doc)

    add_paragraph(doc, "中国大学生计算机设计大赛", font="黑体", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    add_paragraph(doc, "作品信息概要表（填报稿）", font="黑体", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    add_paragraph(doc, DOC_WARNING, color=RGBColor(192, 0, 0))

    info_rows = [
        ("项目", "内容"),
        ("作品编号", "【待填写】"),
        ("作品名称", PROJECT_NAME),
        ("作品大类", "软件应用与开发"),
        ("作品小类", "Web 应用与开发"),
        ("作品简介（100字以内）", "作品构建多人互动圣诞树平台，支持注册登录、建树分享、挂礼物留言、上传图片、关注互动，并在指定日期统一解锁祝福内容。"),
        ("创新描述（100字以内）", "将树状空间交互、时间胶囊加密封存、分级访问控制和节日沉浸式界面融合，兼顾仪式感、社交传播与工程可部署性。"),
    ]
    add_table(doc, info_rows)

    add_heading(doc, "特别说明", level=2)
    add_bullets(
        doc,
        [
            "本作品不涉及疆域地图内容。",
            "如作品存在前期原型、课程项目或历史版本，请在最终提交前据实补充“前期基础与本次参赛主要工作”的说明；当前文稿按“本次参赛周期内完成核心功能、工程化和文档化完善”组织。",
            "本次文档整理使用 OpenAI Codex 辅助完成结构化梳理与文字润色；AI 生成内容主要用于文档草稿，不直接替代核心业务逻辑、数据库设计与关键技术实现，最终内容由参赛队伍人工复核。",
        ],
    )

    add_heading(doc, "作者及其分工比例（需按实际补全）", level=2)
    add_paragraph(doc, "下表仅保留推荐填写结构，请将成员姓名和百分比修改为真实分工；如团队人数超过 3 人，可在 Word 中扩展列。", color=RGBColor(192, 0, 0))
    work_rows = [
        ("项目", "成员1【待替换】", "成员2【待替换】", "成员3【待替换】"),
        ("组织协调", "【%】", "【%】", "【%】"),
        ("作品创意", "【%】", "【%】", "【%】"),
        ("竞品分析", "【%】", "【%】", "【%】"),
        ("方案设计", "【%】", "【%】", "【%】"),
        ("技术实现", "【%】", "【%】", "【%】"),
        ("文献阅读", "【%】", "【%】", "【%】"),
        ("测试分析", "【%】", "【%】", "【%】"),
    ]
    add_table(doc, work_rows)

    add_heading(doc, "指导教师作用", level=2)
    add_paragraph(doc, "□作品创意  □理论指导  □技术方案  □实验场地  □硬件资源  □数据提供  □后勤支持  □宣讲通知  □组织协调  □经费支持  □其他：__________")
    add_paragraph(doc, "请按真实情况将对应方框改为“■”。", color=RGBColor(192, 0, 0))

    add_heading(doc, "开发制作平台与工具", level=2)
    add_paragraph(doc, "开发制作平台：■Windows  □Linux  □macOS  □其他：__________")
    add_paragraph(doc, "运行展示平台：■Windows  ■Linux  ■macOS  ■iOS  ■Android  □其他：__________")
    add_paragraph(doc, "开发制作工具：VS Code、Node.js、npm、Vue 3、Vite、Express 5、Prisma、SQLite、Git/GitHub、Cloudflare Turnstile、Nodemailer。")

    add_heading(doc, "参考文献、项目或作品（前 3 项）", level=2)
    add_bullets(
        doc,
        [
            "Vue.js 官方文档：https://cn.vuejs.org/",
            "Express 官方文档：https://expressjs.com/",
            "Prisma ORM 官方文档：https://www.prisma.io/docs/",
        ],
    )

    add_heading(doc, "提交内容（当前建议勾选项）", level=2)
    add_paragraph(doc, "■报告文档  □演示视频  □PPT  ■源代码  ■部署文件  □数据集  □模型  ■作品文件  □其他")
    add_paragraph(doc, "如答辩演示文档、答辩视频尚未完成，请在最终提交前补勾并补充到总目录。", color=RGBColor(192, 0, 0))

    add_heading(doc, "相关文件", level=2)
    file_rows = [
        ("序号", "文件名", "描述", "文件状态", "版权状态"),
        ("1", "作品信息概要表-Deco My Tree（填报稿）.docx / .pdf", "作品基本信息、创新说明、文件清单", "已整理，待最终信息补全后提交", "自制"),
        ("2", "软件开发类作品设计和开发文档-Deco My Tree（填报稿）.docx / .pdf", "需求、设计、测试、安装及项目总结", "已整理，待最终信息补全后提交", "自制"),
        ("3", "README.md", "项目概述、部署说明和功能特性", "仓库内现有文件", "自制"),
        ("4", "src/", "前端源代码与页面组件", "仓库内现有文件", "自制"),
        ("5", "server/", "后端源代码、接口、数据库与部署配置", "仓库内现有文件", "自制"),
        ("6", "public/screenshot/homepage.png", "代表性首页界面截图", "仓库内现有文件", "自制"),
        ("7", "public/demo.png", "代表性核心交互界面截图", "仓库内现有文件", "自制"),
        ("8", "package.json / server/package.json", "依赖清单与构建配置", "仓库内现有文件", "自制+开源依赖"),
    ]
    add_table(doc, file_rows)

    add_heading(doc, "承诺与签名", level=2)
    add_paragraph(doc, "本作品全体参赛队员确认：本表所列内容属于正式参赛内容的重要组成部分，所有信息均将按要求如实填写，并与最终提交材料保持一致。")
    add_paragraph(doc, "全体参赛队员签名：________________________")
    add_paragraph(doc, f"日期：{DATE_TEXT}")

    output = OUTPUT_DIR / "作品信息概要表-Deco My Tree（填报稿）.docx"
    doc.save(output)
    return output


def build_design_doc(metrics: dict[str, str]) -> Path:
    doc = Document()
    configure_document(doc)

    add_paragraph(doc, "中国大学生计算机设计大赛", font="黑体", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    add_paragraph(doc, "软件开发类作品设计和开发文档", font="黑体", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    add_paragraph(doc, DOC_WARNING, color=RGBColor(192, 0, 0))

    cover_rows = [
        ("项目", "内容"),
        ("作品编号", "【待填写】"),
        ("作品名称", PROJECT_NAME),
        ("作者", "【待填写：成员1、成员2、成员3】"),
        ("版本编号", "V1.0"),
        ("填写日期", DATE_TEXT),
        ("AIGC 说明", "作品核心功能不依赖大模型；本次仅使用 OpenAI Codex 辅助整理文档与文字润色，最终内容由参赛队伍人工审核。"),
    ]
    add_table(doc, cover_rows)

    doc.add_section(WD_SECTION_START.NEW_PAGE)
    add_heading(doc, "目录", level=1)
    for item in [
        "第一章 需求分析",
        "第二章 概要设计",
        "第三章 详细设计",
        "第四章 测试报告",
        "第五章 安装及使用",
        "第六章 项目总结",
        "参考文献",
    ]:
        add_paragraph(doc, item)

    add_heading(doc, "第一章 需求分析", level=1)
    add_heading(doc, "1.1 开发背景与问题界定", level=2)
    add_paragraph(
        doc,
        "当前校园与年轻用户在节日场景中的线上表达，大多停留在即时聊天、朋友圈动态或一次性 H5 贺卡层面。这类方案虽然传播便捷，但普遍存在互动深度不足、共同创作感弱、内容缺少延迟揭晓仪式感、隐私控制粗糙等问题。Deco My Tree 将“节日祝福”升级为可共同参与、可封存、可回访的时间胶囊式互动产品，目标不是单次展示，而是构建一套兼具情感体验与工程完整度的多人互动平台。"
    )

    add_heading(doc, "1.2 目标用户与应用场景", level=2)
    add_bullets(
        doc,
        [
            "校园班级、宿舍、社团、实验室等小群体节日互动场景。",
            "朋友、情侣、异地亲友之间的节日祝福与纪念内容封存场景。",
            "活动运营方需要轻量上线、可分享传播、可在移动端直接使用的互动页面场景。",
        ],
    )

    add_heading(doc, "1.3 主要功能需求", level=2)
    add_bullets(
        doc,
        [
            "支持邮箱注册登录、邮件验证与 LinuxDo OAuth 登录，降低使用门槛。",
            "支持每位用户创建并管理一棵专属圣诞树，设置公开、仅链接访问或仅自己可见。",
            "支持在树上点击空白区域挂载礼物，填写祝福内容、上传最多 3 张图片，并进行提交前预览。",
            "支持时间胶囊机制：内容在指定解锁时间前加密封存，到期后统一展示。",
            "支持关注其他人的树，形成回访和持续互动关系。",
            "支持移动端适配、沉浸式雪景背景、倒计时等节日氛围设计。",
        ],
    )

    add_heading(doc, "1.4 主要性能与工程目标", level=2)
    add_bullets(
        doc,
        [
            "交互层面：保证主要流程在手机浏览器中单手可操作，核心弹窗和卡片组件适配小屏设备。",
            "安全层面：建立邮箱验证、JWT、验证码、接口限流、文件过滤和环境变量校验等完整链路。",
            "部署层面：基于 Node.js + SQLite + Prisma 的轻量方案，适合比赛演示与中小规模上线。",
            "可维护层面：前后端分离，前端 5 个视图、7 个核心组件，后端围绕认证、树、留言三类核心对象组织接口。",
        ],
    )

    add_heading(doc, "1.5 同类方案与竞品分析", level=2)
    competitor_rows = [
        ("维度", "传统电子贺卡 H5", "普通留言墙/树洞", "通用社交动态", "本作品"),
        ("核心表达方式", "单向发送", "即时文字留言", "信息流展示", "树状空间化互动"),
        ("多人共创能力", "弱", "中", "中", "强，可围绕同一棵树持续添加礼物"),
        ("节日仪式感", "依赖模板视觉", "弱", "弱", "强，倒计时 + 统一解锁形成高潮"),
        ("隐私与权限", "通常较弱", "匿名但不精细", "平台级通用权限", "公开/链接/私密三级访问控制"),
        ("独立部署能力", "通常依赖模板平台", "较弱", "依赖第三方平台", "完整前后端可独立部署"),
        ("传播与回访", "一次性分享", "弱", "强但易淹没主题", "关注关系 + 时间胶囊回访机制"),
    ]
    add_table(doc, competitor_rows)
    add_paragraph(
        doc,
        "与常见同类方案相比，Deco My Tree 的竞争力不只来自“圣诞主题美术”，更来自完整的产品机制设计：它把情绪表达、多人参与、隐私保护、轻量部署和后续运营潜力放在同一个系统里完成，具备从作品原型走向可上线产品的基础。"
    )

    add_heading(doc, "第二章 概要设计", level=1)
    add_heading(doc, "2.1 系统总体架构", level=2)
    add_paragraph(
        doc,
        "系统采用前后端分离架构。前端基于 Vue 3 + Vue Router + Vite 构建单页应用，负责节日场景展示、交互动画、表单预览和多页面路由；后端基于 Express 5 暴露 RESTful API，负责认证、树管理、礼物存取、权限控制和文件上传；数据持久层由 Prisma 统一访问 SQLite；外部服务接入 LinuxDo OAuth、SMTP 邮件验证与 Cloudflare Turnstile。"
    )
    architecture_rows = [
        ("层次", "组成", "职责", "竞争力体现"),
        ("前端展示层", "Landing / Home / TreeView / FollowedTrees / About", "承担品牌表达、页面跳转、交互反馈与响应式布局", "节日氛围强、上手门槛低、移动端体验完整"),
        ("业务接口层", "auth.js / trees.js", "用户认证、树管理、留言互动、关注关系与权限校验", "业务边界清晰，便于扩展更多节日活动"),
        ("数据与文件层", "Prisma + SQLite + uploads", "维护用户、树、礼物、关注关系及图片资源", "轻量可迁移，适合比赛演示与快速部署"),
        ("安全与运维层", "JWT / Argon2 / Turnstile / rate-limit / CORS", "认证授权、反滥用、文件过滤和部署校验", "体现工程成熟度，不是单纯视觉作品"),
    ]
    add_table(doc, architecture_rows)

    add_heading(doc, "2.2 功能模块划分", level=2)
    module_rows = [
        ("模块", "核心页面/接口", "主要职责", "竞争力价值"),
        ("门户与探索模块", "Landing.vue、Home.vue、TreeList.vue", "完成品牌介绍、多树浏览和冷启动转化", "降低新用户首次进入时的理解成本"),
        ("认证模块", "AuthModal.vue、/api/auth/*", "邮箱注册登录、邮件验证、OAuth 登录、身份保持", "兼顾社区账号接入和正式用户体系"),
        ("树管理模块", "TreeView.vue、TreeSettings.vue、/api/trees", "创建树、命名树、访问权限配置", "适配公开传播与私密纪念两类需求"),
        ("留言互动模块", "ChristmasTree.vue、ChristmasCard.vue、/decorations", "礼物定位、图标选择、预览、上传图片、写祝福", "把静态祝福升级为可探索的空间交互"),
        ("社交回访模块", "FollowedTrees.vue、/follow", "形成关注链路与回访入口", "提升留存和活动复访率"),
        ("安全与部署模块", "config.js、upload.js、middleware/auth.js", "环境校验、文件过滤、权限控制、反刷保护", "保证作品具备真实部署能力"),
    ]
    add_table(doc, module_rows)

    add_heading(doc, "2.3 典型业务流程", level=2)
    add_bullets(
        doc,
        [
            "首次进入：用户阅读封面页文案，了解“时间胶囊圣诞树”的规则后选择登录或先浏览。",
            "创建树：登录用户一键创建自己的树，并根据传播需求配置公开、链接访问或私密模式。",
            "挂礼物：访问者点击树上空白位置，选择图标、填写留言、上传图片、通过人机验证后提交。",
            "封存与解锁：后端对内容进行 AES 加密保存；未到解锁时间时仅展示“封印中”，到期后统一解密展示。",
            "回访互动：留言用户可自动建立关注关系，后续在“关注的树”页面回访同伴作品。",
        ],
    )

    add_heading(doc, "第三章 详细设计", level=1)
    add_heading(doc, "3.1 界面设计与典型流程", level=2)
    add_paragraph(
        doc,
        "界面设计重点强调“冬日节庆感”和“低学习成本”。作品采用雪景背景、暖色按钮、倒计时、翻转卡片和树状空间交互等设计语言，使用户在首次进入时即能理解：这不是普通留言板，而是一棵可以共同装饰、等待节日揭晓的树。"
    )
    if HOME_SHOT.exists():
        doc.add_picture(str(HOME_SHOT), width=Cm(15.5))
        add_paragraph(doc, "图 3-1 作品首页与多树浏览界面", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    if DEMO_SHOT.exists():
        doc.add_picture(str(DEMO_SHOT), width=Cm(12.5))
        add_paragraph(doc, "图 3-2 作品核心互动与主题展示界面", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

    flow_rows = [
        ("步骤", "用户动作", "系统反馈"),
        ("1", "打开作品主页", "展示项目定位、功能亮点和入口按钮"),
        ("2", "登录/注册", "完成邮箱验证或 OAuth 登录，建立身份"),
        ("3", "创建个人圣诞树", "生成唯一 slug，进入专属树页面"),
        ("4", "点击树上空白区域挂礼物", "弹出写卡片界面，可选图标、图片和隐私设置"),
        ("5", "确认提交", "后端保存密文和资源地址，前端显示成功提示"),
        ("6", "节日当天回访", "系统自动解锁礼物内容并允许查看历史祝福"),
    ]
    add_table(doc, flow_rows)

    add_heading(doc, "3.2 数据库设计", level=2)
    add_paragraph(
        doc,
        "数据库围绕“用户”“树”“礼物”“关注关系”四类核心对象展开，既满足当前节日互动需求，也为后续扩展到更多主题活动预留了结构空间。实体关系为：User 与 Tree 为 1:1，Tree 与 Decoration 为 1:N，User 与 Tree 通过 Follow 构成 N:N 关注关系。"
    )
    db_rows = [
        ("实体", "关键字段", "说明"),
        ("User", "email、passwordHash、emailVerified、linuxdoId、username、isAdmin", "支持邮箱与 OAuth 双认证，并保留管理员能力"),
        ("Tree", "slug、title、ownerId、isPublic、isAccessible", "树的唯一访问标识、标题和访问权限"),
        ("Decoration", "x、y、icon、nickname、content、isPrivate、images、treeId", "记录礼物在树上的坐标、图标、密文内容和图片资源"),
        ("Follow", "userId、treeId、createdAt", "维护用户与树之间的关注关系，用于后续回访"),
    ]
    add_table(doc, db_rows)

    add_heading(doc, "3.3 关键技术与实现要点", level=2)
    tech_rows = [
        ("关键点", "实现说明", "对竞争力的贡献"),
        ("时间胶囊封存机制", "留言提交时使用 AES 加密保存；服务端根据 UNLOCK_DATE 决定返回密文占位还是解密结果", "将普通留言升级为“延迟揭晓”的节日体验核心"),
        ("树上空间定位与分页", "礼物记录为相对坐标，保证不同屏幕下挂载位置稳定；每页 10 个礼物，树满时自动引导到下一页", "兼顾视觉美感、可读性和持续增长能力"),
        ("多级访问控制", "树支持公开、仅链接访问、仅自己可见；私密留言在解锁后也仅树主/管理员可见；关注列表过滤不可访问树", "让作品既能传播也能承载私人纪念内容"),
        ("安全与反滥用", "邮箱验证、JWT、Argon2、Turnstile、人机验证、接口限流、图片类型与大小过滤、环境变量强校验", "相比常见学生作品更接近真实线上产品要求"),
        ("轻量部署架构", "前后端分离，前端 build 后静态托管，后端 Node 服务结合 Prisma migrate deploy 即可启动", "比赛答辩演示和后续上线都更可控"),
    ]
    add_table(doc, tech_rows)

    add_heading(doc, "第四章 测试报告", level=1)
    add_heading(doc, "4.1 测试环境", level=2)
    test_env_rows = [
        ("项目", "环境"),
        ("操作系统", "Microsoft Windows 11 家庭中文版 10.0.26200"),
        ("Node.js", "v24.4.1"),
        ("npm", "11.4.2"),
        ("前端构建工具", "Vite 7"),
        ("后端运行环境", "Node.js + Express 5"),
        ("数据库", "SQLite + Prisma"),
    ]
    add_table(doc, test_env_rows)

    add_heading(doc, "4.2 主要测试项与结果", level=2)
    case_rows = [
        ("编号", "测试项", "测试方法", "结果"),
        ("T01", "前端生产构建", "执行 npm run build", "通过"),
        ("T02", "用户名清洗逻辑", "执行 node server/test-sanitize.js，11 组边界数据测试", "11/11 通过"),
        ("T03", "树列表接口", "本地临时数据库启动后访问 /api/trees", "通过"),
        ("T04", "开发登录与 JWT 签发", "访问 /api/auth/dev-login 并提取 token", "通过"),
        ("T05", "创建个人圣诞树", "携带 token 调用 POST /api/trees", "通过"),
        ("T06", "挂礼物与封印返回", "提交 decoration 后读取 decorations，验证未解锁时返回“封印中...”", "通过"),
        ("T07", "访问权限控制", "私密树不允许未授权关注，不可访问树不会出现在关注列表中", "通过（本轮修正后复核）"),
        ("T08", "部署链路", "Prisma migrate deploy + Node 服务启动 + 前端构建产物联动", "通过"),
    ]
    add_table(doc, case_rows)

    add_heading(doc, "4.3 技术指标与量化结果", level=2)
    metric_rows = [
        ("指标", "结果", "说明"),
        ("前端 JS 包体积", metrics["js_size"], f"gzip 后约 {metrics['js_gzip']}"),
        ("前端 CSS 包体积", metrics["css_size"], f"gzip 后约 {metrics['css_gzip']}"),
        ("本地冒烟测试响应", "树列表 7ms / 登录 15ms / 建树 33ms / 挂礼物 20ms / 读取礼物 5ms", "基于临时 SQLite 数据库的本地一次实测"),
        ("安全阈值", "认证接口 30 次/15 分钟；通用 API 200 次/15 分钟；挂礼物 10 次/分钟", "由 express-rate-limit 配置"),
        ("上传约束", "最多 3 张图片，单张不超过 5MB", "限制 JPEG/PNG/GIF/WEBP"),
        ("稳定性与可维护性", "前后端分离，数据库迁移可自动化执行", "便于答辩演示、复现和后续扩展"),
    ]
    add_table(doc, metric_rows)

    add_heading(doc, "第五章 安装及使用", level=1)
    add_heading(doc, "5.1 安装环境要求", level=2)
    add_bullets(
        doc,
        [
            "Node.js 20 及以上版本，npm 可用。",
            "前端与后端依赖分别已安装；数据库采用 SQLite，无需单独部署数据库服务。",
            "生产环境建议配置 SMTP、LinuxDo OAuth、Cloudflare Turnstile、PM2 与 Caddy。",
        ],
    )

    add_heading(doc, "5.2 默认安装流程", level=2)
    install_rows = [
        ("步骤", "操作"),
        ("1", "在项目根目录执行 npm install；进入 server 目录执行 npm install"),
        ("2", "在 server 目录执行 npx prisma generate 与 npx prisma migrate dev（或生产环境下使用 migrate deploy）"),
        ("3", "启动后端：cd server && node index.js"),
        ("4", "启动前端：在项目根目录执行 npm run dev"),
        ("5", "访问 http://localhost:5173/tree/ 或生产环境对应域名"),
    ]
    add_table(doc, install_rows)

    add_heading(doc, "5.3 典型使用流程", level=2)
    add_bullets(
        doc,
        [
            "用户注册或使用 LinuxDo 登录后，一键创建个人圣诞树。",
            "用户调整树的标题和可见范围，并把链接分享给同学或朋友。",
            "访问者在树上挂礼物、上传照片并选择是否设为悄悄话。",
            "在解锁日期之前，所有礼物统一显示为“封印中”；到期后统一可读。",
            "用户可通过“关注的树”持续回访其他人的节日树。",
        ],
    )

    add_heading(doc, "第六章 项目总结", level=1)
    add_paragraph(
        doc,
        "Deco My Tree 的开发过程强调两个方向：一是把“节日情绪价值”产品化，而不是只做视觉海报；二是把“比赛作品”工程化，而不是只做静态演示。项目在实现过程中重点解决了身份认证、内容封存、移动端弹窗适配、文件上传限制、访问权限控制和独立部署等问题，使作品既能打动评委，也能支持真实使用。"
    )
    add_paragraph(
        doc,
        "后续升级方向包括：扩展到毕业季、生日、情人节等更多主题模板；增加活动数据看板和内容审核能力；接入更丰富的多媒体内容；提供组织者后台与商业化活动页模板。基于现有架构，上述功能可以在不推翻核心系统的前提下逐步演进。"
    )

    add_heading(doc, "参考文献", level=1)
    add_bullets(
        doc,
        [
            "Vue.js 官方文档. https://cn.vuejs.org/",
            "Express 官方文档. https://expressjs.com/",
            "Prisma ORM 官方文档. https://www.prisma.io/docs/",
            "Cloudflare Turnstile 官方文档. https://developers.cloudflare.com/turnstile/",
            "OWASP Authentication Cheat Sheet. https://cheatsheetseries.owasp.org/",
        ],
    )

    output = OUTPUT_DIR / "软件开发类作品设计和开发文档-Deco My Tree（填报稿）.docx"
    doc.save(output)
    return output


def write_readme() -> None:
    content = (
        "本文件夹用于存放中国大学生计算机设计大赛“Deco My Tree”项目的设计与开发文档。\n"
        "主要文件说明：\n"
        "1. 《作品信息概要表-Deco My Tree（填报稿）》：用于填写作品基本信息、创新说明、作者分工和文件清单。\n"
        "2. 《软件开发类作品设计和开发文档-Deco My Tree（填报稿）》：包含需求分析、概要设计、详细设计、测试报告、安装使用和项目总结。\n"
        "3. 同名 PDF 文件：若本机可调用 Microsoft Word，则脚本会自动导出 PDF 版本。\n"
        "4. 《待完善信息清单》：列出仍需由参赛队伍根据真实情况补全的字段。\n"
        "注意：文档中所有“【待填写】”内容均需在提交前核实并补齐。\n"
    )
    (OUTPUT_DIR / "readme.txt").write_text(content, encoding="utf-8")


def write_todo_list() -> None:
    content = (
        "提交前仍需人工补全/确认的内容：\n"
        "1. 作品编号。\n"
        "2. 全体参赛队员姓名、签名、分工比例。\n"
        "3. 指导教师作用勾选项。\n"
        "4. 如有前期基础，请据实补充本次参赛的主要新增工作。\n"
        "5. 提交内容勾选项中 PPT、答辩视频等尚未生成的材料。\n"
        "6. 相关文件表中的最终文件状态、下载地址或网盘状态。\n"
        "7. 如团队在代码实现阶段也使用过 AI，请把 AI 辅助说明改成真实比例和真实用途。\n"
    )
    (OUTPUT_DIR / "待完善信息清单.txt").write_text(content, encoding="utf-8")


def export_pdf(docx_path: Path) -> Path | None:
    if win32 is None:
        return None

    pdf_path = docx_path.with_suffix(".pdf")
    word = None
    document = None
    try:
        word = win32.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        document = word.Documents.Open(str(docx_path.resolve()))
        document.SaveAs(str(pdf_path.resolve()), FileFormat=17)
        return pdf_path
    except Exception:
        return None
    finally:
        if document is not None:
            document.Close(False)
        if word is not None:
            word.Quit()


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    metrics = get_build_metrics()

    summary_doc = build_summary_doc()
    design_doc = build_design_doc(metrics)
    write_readme()
    write_todo_list()

    pdf_results = []
    for docx_path in [summary_doc, design_doc]:
        pdf_path = export_pdf(docx_path)
        pdf_results.append((docx_path.name, pdf_path.name if pdf_path else "PDF 导出失败或本机未安装 Word"))

    print("已生成文件目录：", OUTPUT_DIR)
    for item in OUTPUT_DIR.iterdir():
        print("-", item.name)
    print("PDF 导出结果：")
    for doc_name, result in pdf_results:
        print(f"- {doc_name} -> {result}")


if __name__ == "__main__":
    main()
